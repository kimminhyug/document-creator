import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile
from openpyxl import load_workbook
from docx import Document
import creator
from output_adapters import render_html, render_xlsx

class OutputTests(unittest.TestCase):
    def test_docx_company_fonts_override_template_themes(self):
        from docx.oxml.ns import qn
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'fonts.docx'
            creator.render_docx(path,self.content,self.project,self.theme)
            doc=Document(path)
            for name in ('Normal','Title','Heading 1','Heading 2','Heading 3','Heading 4','Header','Footer'):
                fonts=doc.styles[name]._element.rPr.rFonts
                for slot in ('ascii','hAnsi','eastAsia','cs'):
                    self.assertEqual(fonts.get(qn('w:'+slot)),self.theme['font_family'])
                self.assertFalse(any(key.lower().endswith('theme') for key in fonts.attrib))

    def test_docx_table_pagination_and_width_policy(self):
        from docx.oxml.ns import qn
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'tables.docx'
            for keep,mode in ((True,'content'),(False,'equal')):
                self.theme['table']={'keep_rows_together':keep,'column_width_mode':mode}
                creator.render_docx(path,self.content,self.project,self.theme)
                for table in Document(path).tables:
                    self.assertFalse(table.autofit)
                    for row in table.rows:
                        for column,cell in zip(table.columns,row.cells):
                            self.assertAlmostEqual(column.width.pt,cell.width.pt,places=1)
                    self.assertIsNotNone(table.rows[0]._tr.trPr.find(qn('w:tblHeader')))
                    for row in table.rows:
                        self.assertEqual(row._tr.trPr is not None and row._tr.trPr.find(qn('w:cantSplit')) is not None,keep)
            self.theme['table']['keep_rows_together']=1
            with self.assertRaises(ValueError):creator.validate_theme(self.theme)

    def test_docx_title_words_and_short_sections_preserve_text(self):
        from docx.oxml.ns import qn
        self.content['title']='설비 모니터링 테스트 제품 기능 명세서'
        self.content['sections']=[{'title':'확인 범위','status':'confirmed','blocks':[
            {'type':'paragraph','text':'첫 번째 짧은 문단.'},
            {'type':'paragraph','text':'두 번째 짧은 문단.'}]}]
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'title.docx';creator.render_docx(path,self.content,self.project,self.theme)
            doc=Document(path);title=next(p for p in doc.paragraphs if p.style.name=='Title')
            self.assertEqual(title.text.split(),self.content['title'].split())
            self.assertIn('명세서',title.text.split('\n')[-1])
            self.assertIsNone(doc.styles['Title']._element.pPr.find(qn('w:pBdr')))
            first=next(p for p in doc.paragraphs if p.text=='첫 번째 짧은 문단.')
            self.assertTrue(first.paragraph_format.keep_with_next)
            last=next(p for p in doc.paragraphs if p.text=='두 번째 짧은 문단.')
            self.assertIsNot(last.paragraph_format.keep_with_next,True)

    def test_docx_mixed_section_does_not_chain_all_paragraphs(self):
        self.content['sections'][0]['blocks'].insert(0, {'type':'paragraph','text':'표 앞의 독립 문단'})
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'mixed.docx'
            creator.render_docx(path,self.content,self.project,self.theme)
            paragraph=next(p for p in Document(path).paragraphs if p.text=='표 앞의 독립 문단')
            self.assertIsNot(paragraph.paragraph_format.keep_with_next,True)

    def test_docx_language_is_explicit_and_configurable(self):
        from docx.oxml.ns import qn
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'language.docx'
            for language in ('ko-KR','en-US','zh-Hant'):
                if language!='ko-KR':self.project['language']=language
                creator.render_docx(path,self.content,self.project,self.theme)
                doc=Document(path)
                self.assertEqual(doc.core_properties.language,language)
                for name in ('Normal','Title','Heading 1','Heading 2','Header','Footer'):
                    element=doc.styles[name]._element.rPr.find(qn('w:lang'))
                    self.assertEqual(element.get(qn('w:eastAsia')),language)
                    self.assertEqual(element.get(qn('w:val')),'en-US' if language.startswith(('ko','zh')) else language)
            self.project['language']='not a language'
            with self.assertRaisesRegex(ValueError,'project.language'):
                creator.validate(self.content,{'id':self.content['type']},self.project,self.theme)

    def test_docx_word_layout_preserves_tokens_and_description_space(self):
        from docx_layout import wrap_words,column_widths
        value='조회 기간은 최대 31일입니다. 검색어는 최대 100자입니다. 00:00~11:00으로 조회합니다.'
        wrapped=wrap_words(value,10.5,100)
        self.assertEqual(wrapped.split(),value.split())
        for token in ('31일입니다.','100자입니다.','00:00~11:00으로'):
            self.assertIn(token,wrapped)
        block={'headers':['객체','사용 항목','용도'],'rows':[['documentation.equipment_state_history','equipment_id,state,start_at,end_at','기준 기간의 로컬 테스트 정답표']]}
        widths=column_widths(block,self.theme,470)
        self.assertAlmostEqual(sum(widths),470)
        self.assertGreaterEqual(widths[2],72)

    def test_public_web_address_is_clickable_without_losing_printed_text(self):
        address='https://example.test/path?a=1&b=2'
        self.content['web_links']=[address]
        self.content['sections'][0]['blocks'].append({'type':'paragraph','text':'접속 ('+address+')'})
        with tempfile.TemporaryDirectory() as temp:
            folder=Path(temp)
            render_html(folder/'doc.html',self.content,self.project,self.theme)
            self.assertIn('href="https://example.test/path?a=1&amp;b=2"',(folder/'doc.html').read_text('utf-8'))
            creator.render_docx(folder/'doc.docx',self.content,self.project,self.theme)
            with ZipFile(folder/'doc.docx') as z:
                self.assertIn('w:hyperlink',z.read('word/document.xml').decode())
                self.assertIn('Target="https://example.test/path?a=1&amp;b=2"',z.read('word/_rels/document.xml.rels').decode())
            self.assertIn(address,'\n'.join(p.text for p in Document(folder/'doc.docx').paragraphs))
            render_xlsx(folder/'doc.xlsx',self.content,self.project,self.theme)
            wb=load_workbook(folder/'doc.xlsx')
            links=[c.hyperlink.target for ws in wb for row in ws for c in row if c.hyperlink]
            self.assertIn(address,links);wb.close()

    def setUp(self):
        self.content=creator.read_json(creator.ROOT/'examples/demo/content/table-spec.json')
        self.project=creator.read_json(creator.ROOT/'examples/demo/content/project.json')
        self.theme=creator.read_json(creator.ROOT/'themes/blue-office.json')
        self.theme['copyright']={'enabled':True,'title':'저작권 안내','text':'Copyright {organization}'}
        self.theme['styles']={'heading_2':{'size_pt':19,'bold':False},'table_header':{'size_pt':13,'bold':False}}
        self.content['sections'][0]['blocks'].extend([{'type':'heading','level':2,'text':'상세 설명'}, {'type':'paragraph','text':'<script>alert(1)</script>'}, {'type':'table','headers':['값'],'rows':[['=HYPERLINK("https://example.com")'],['+1'],['@SUM(A1)']]}])
    def test_html_escapes_and_does_not_publish_internal_evidence(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'doc.html';render_html(path,self.content,self.project,self.theme);data=path.read_text('utf8')
            self.assertIn('&lt;script&gt;',data);self.assertNotIn('<script>',data);self.assertIn('Content-Security-Policy',data)
            self.assertIn('저작권 안내',data)
            for source in self.content['sources']:self.assertNotIn(source['locator'],data);self.assertNotIn(source['id'],data)
            self.assertNotIn(self.content['id'],data)
    def test_xlsx_formula_strings_and_header_style(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'doc.xlsx';render_xlsx(path,self.content,self.project,self.theme);wb=load_workbook(path)
            cells=[c for ws in wb for row in ws for c in row if c.value is not None]
            for cell in cells:self.assertNotEqual(cell.data_type,'f')
            header=next(c for c in cells if c.value=='값');self.assertEqual(header.font.sz,13);self.assertFalse(header.font.bold)
            self.assertEqual(header.fill.fgColor.rgb[-6:],self.theme['colors']['table_fill'][1:])
            values='\n'.join(str(c.value) for c in cells)
            self.assertNotIn(self.content['sources'][0]['id'],values);self.assertIn('=HYPERLINK',values)
            wb.close()
    def test_docx_styles_and_no_internal_ids(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'doc.docx';creator.render_docx(path,self.content,self.project,self.theme);doc=Document(path)
            self.assertEqual(doc.styles['Heading 2'].font.size.pt,19);self.assertFalse(doc.styles['Heading 2'].font.bold)
            with ZipFile(path) as z:data=z.read('word/document.xml').decode()
            self.assertIn('저작권 안내',data);self.assertNotIn(self.content['id'],data)
            for source in self.content['sources']:self.assertNotIn(source['id'],data);self.assertNotIn(source['locator'],data)
    def test_local_image_embedded_without_local_path(self):
        from PIL import Image
        with tempfile.TemporaryDirectory() as tmp:
            source=Path(tmp)/'private-image.png';Image.new('RGB',(30,20),'blue').save(source)
            self.content['sections'][0]['blocks'].append({'type':'image','path':str(source),'alt':'Overview'})
            path=Path(tmp)/'doc.html';render_html(path,self.content,self.project,self.theme)
            html=path.read_text('utf8');self.assertIn('data:image/png;base64,',html);self.assertNotIn(str(source),html)
            xlsx=Path(tmp)/'doc.xlsx';render_xlsx(xlsx,self.content,self.project,self.theme)
            with ZipFile(xlsx) as z:self.assertIn('xl/media/image1.png',z.namelist())

    def test_invalid_css_and_public_url_rejected(self):
        self.theme['styles']['body']={'font_family':'</style><script>'}
        with self.assertRaises(ValueError):creator.validate_theme(self.theme)
        from output_adapters import public_references
        with self.assertRaises(ValueError):public_references({'public_references':[{'title':'bad','url':'javascript:alert(1)'}]})
    def test_build_html_xlsx_atomic_manifest(self):
        with tempfile.TemporaryDirectory() as tmp:
            out=creator.build(creator.ROOT/'examples/demo/content/table-spec.json',creator.ROOT/'templates/table_spec.json',creator.ROOT/'examples/demo/content/project.json',creator.ROOT/'themes/blue-office.json',None,tmp,['html','xlsx'])
            manifest=creator.read_json(out/'manifest.json');self.assertEqual(set(manifest['outputs']),{'document.html','document.xlsx'})

if __name__=='__main__': unittest.main()
