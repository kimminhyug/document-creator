"""Human-facing HTML and spreadsheet renderers; provenance stays in the run manifest."""
from __future__ import annotations
import base64
from html import escape
from pathlib import Path
import re
from urllib.parse import urlsplit
from inline_links import safe_url_markup, url_parts


def style_token(theme, role):
    sizes = {'title': theme['title_pt'], 'heading_1': theme['heading_pt'], 'heading_2': theme['body_pt']+2, 'heading_3': theme['body_pt']+1, 'heading_4': theme['body_pt'], 'body': theme['body_pt'], 'table_header': theme['table_pt'], 'table_body': theme['table_pt'], 'header': 9, 'footer': 9}
    before = {'heading_1': 20, 'heading_2': 16, 'heading_3': 14, 'heading_4': 12}.get(role, 0)
    after = 10 if role.startswith('heading_') else {'title': 14, 'body': 9}.get(role, 0)
    return {'font_family': theme['font_family'], 'size_pt': sizes[role], 'bold': role.startswith('heading_') or role in ('title', 'table_header'), 'space_before_pt': before, 'space_after_pt': after, **theme.get('styles', {}).get(role, {})}


def version_text(content, theme):
    label = theme.get('cover', {}).get('version_label', '버전')
    return (label + ' ' + content['version']).strip()


def public_references(content):
    refs = content.get('public_references', [])
    if not isinstance(refs, list):
        raise ValueError('public_references must be a list')
    for ref in refs:
        if not isinstance(ref, dict) or not isinstance(ref.get('title'), str) or not ref['title'].strip():
            raise ValueError('public reference requires human-readable title')
        if ref.get('url'):
            parsed = urlsplit(ref['url'])
            if parsed.scheme not in ('http', 'https') or not parsed.hostname or parsed.username or parsed.password:
                raise ValueError('public reference URL must be HTTP(S), without credentials')
    return refs


def render_html(path, content, project, theme):
    from creator import section_blocks, interpolate
    def e(value): return escape(str(value), quote=True)
    def prose(value): return safe_url_markup(str(value), content.get('web_links', []))
    def fmt(value): return interpolate(value, content, project)
    rules = []
    for role, selector in [('title', '.title'), ('heading_1','h1'), ('heading_2','h2'), ('heading_3','h3'), ('heading_4','h4'), ('body','body'), ('table_header','th'), ('table_body','td'), ('header','header'), ('footer','footer')]:
        token = style_token(theme, role)
        # Quote a validated family name; no raw CSS from product content.
        family = token['font_family'].replace('\\','\\\\').replace('"','\\"')
        rules.append(f'{selector}{{font-family:"{family}",sans-serif;font-size:{token["size_pt"]}pt;font-weight:{700 if token["bold"] else 400};margin-top:{token["space_before_pt"]}pt;margin-bottom:{token["space_after_pt"]}pt}}')
    c = theme['colors']; table = theme.get('table', {})
    css = '\n'.join(rules) + f'\nbody{{margin:0;background:{c["background"]};color:{c["text"]};line-height:{theme["line_spacing"]}}}main{{max-width:980px;margin:auto;padding:40px}}header,footer{{color:{c["muted"]};padding:18px 0}}table{{border-collapse:collapse;width:100%;margin:16px 0}}th,td{{text-align:left;vertical-align:top;border-bottom:{table.get("border_pt",.35)}pt solid {c["border"]};padding:{table.get("padding_y_pt",7)}pt {table.get("padding_x_pt",7)}pt;overflow-wrap:anywhere}}th{{background:{c["table_fill"]};color:{c["table_text"]}}}img{{max-width:100%;height:auto}}pre{{white-space:pre-wrap;overflow-wrap:anywhere}}p{{white-space:pre-wrap}}section{{margin-bottom:32px}}a{{color:{c["accent"]}}}@media print{{main{{padding:0}}.front{{break-after:page}}thead{{display:table-header-group}}}}'
    html = ['<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">', '<meta http-equiv="Content-Security-Policy" content="default-src &#39;none&#39;; img-src data:; style-src &#39;unsafe-inline&#39;; base-uri &#39;none&#39;; form-action &#39;none&#39;">', '<title>'+e(content['title'])+'</title><style>'+css+'</style></head><body><main>']
    html.append('<header>'+e(fmt(theme['header']['left']))+' · '+e(fmt(theme['header']['right']))+'</header>')
    if theme.get('cover',{}).get('enabled',True):
        html.append('<section class="front"><p>'+e(project['organization'])+'</p><h1 class="title">'+e(fmt(theme.get('cover',{}).get('title','{title}')) )+'</h1><p>'+e(content['summary'])+'</p><p>'+e(version_text(content, theme))+' · '+e(project['classification'])+'</p>'+('<p>가상 예제 데이터로 작성된 문서입니다.</p>' if content.get('example') else '')+'</section>')
    if theme.get('copyright',{}).get('enabled',False):
        html.append('<section class="front"><h1>'+e(fmt(theme['copyright'].get('title','저작권')))+'</h1><p>'+e(fmt(theme['copyright'].get('text','')))+'</p></section>')
    if theme.get('contents',{}).get('enabled',True):
        html.append('<nav class="front" aria-label="목차"><h1>'+e(theme.get('contents',{}).get('title','문서 구성'))+'</h1><ol>')
        for n,section in enumerate(content['sections'],1): html.append(f'<li><a href="#section-{n}">'+e(section['title'])+'</a></li>')
        html.append('</ol></nav>')
    for n,section in enumerate(content['sections'],1):
        html.append(f'<section id="section-{n}"><h1>{n} '+e(section['title'])+'</h1>')
        for b in section_blocks(section):
            k=b['type']
            if k=='table':
                html.append('<table><thead><tr>'+''.join('<th scope="col">'+prose(h)+'</th>' for h in b['headers'])+'</tr></thead><tbody>')
                html.extend('<tr>'+''.join('<td>'+prose(v)+'</td>' for v in row)+'</tr>' for row in b['rows'])
                html.append('</tbody></table>')
            elif k=='image':
                p=Path(b['path']); mime='image/png' if p.suffix.lower()=='.png' else 'image/jpeg'
                html.append('<figure><img src="data:'+mime+';base64,'+base64.b64encode(p.read_bytes()).decode('ascii')+'" alt="'+e(b['alt'])+'"><figcaption>'+e(b['alt'])+'</figcaption></figure>')
            elif k=='heading': html.append(f'<h{b["level"]}>'+e(b['text'])+f'</h{b["level"]}>')
            elif k=='code': html.append('<pre><code>'+e(b['text'])+'</code></pre>')
            elif k=='note':
                note=theme['notes'][b['role']]
                html.append('<aside style="padding:12px;background:'+note['fill']+';border-left:4px solid '+note['marker']+'"><strong>'+e(note['label'])+'</strong><p>'+prose(b['text'])+'</p></aside>')
            else: html.append('<p>'+prose(b['text'])+'</p>')
        html.append('</section>')
    if public_references(content):
        html.append('<h1>참고 자료</h1><ul>')
        for ref in public_references(content):
            html.append('<li>'+ ('<a rel="noopener noreferrer" href="'+e(ref['url'])+'">'+e(ref['title'])+'</a>' if ref.get('url') else e(ref['title']))+'</li>')
        html.append('</ul>')
    if theme['ending'].get('enabled',True): html.append('<section class="front"><h1>'+e(fmt(theme['ending']['title']))+'</h1><p>'+e(fmt(theme['ending']['text']))+'</p></section>')
    html.append('<footer>'+e(fmt(theme['footer']['left']))+'</footer></main></body></html>')
    Path(path).write_text('\n'.join(html),encoding='utf-8')


def render_xlsx(path, content, project, theme):
    from creator import section_blocks, interpolate
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.drawing.image import Image
    wb=Workbook(); wb.remove(wb.active)
    def fmt(v): return interpolate(v,content,project)
    def font(role, color=None):
        token=style_token(theme,role)
        return Font(name=token['font_family'],size=token['size_pt'],bold=token['bold'],color=color or theme['colors']['text'][1:])
    def sheet(title):
        title=re.sub(r'[\\/*?:\[\]]',' ',title).strip()[:31] or '문서'
        candidate=title; n=2
        while candidate.lower() in [s.lower() for s in wb.sheetnames]: candidate=title[:26]+f' ({n})'; n+=1
        ws=wb.create_sheet(candidate);ws.sheet_view.showGridLines=False;ws.freeze_panes='A3'
        for col in 'ABCDEF': ws.column_dimensions[col].width=25
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
        # Ampersands are Excel print-field escapes, so quote all content ampersands.
        ws.oddHeader.left.text=fmt(theme['header']['left']).replace('&','&&')
        ws.oddHeader.right.text=fmt(theme['header']['right']).replace('&','&&')
        ws.oddFooter.left.text=fmt(theme['footer']['left']).replace('&','&&'); ws.oddFooter.right.text='&P / &N'
        for area, role in [(ws.oddHeader,'header'),(ws.oddFooter,'footer')]:
            token=style_token(theme,role)
            for part in (area.left,area.center,area.right): part.size=token['size_pt'];part.font=token['font_family'] + (',Bold' if token['bold'] else ',Regular')
        return ws
    def cell(ws,row,col,value,role='body'):
        value=str(value)
        if len(value)>32767: raise ValueError('Excel cell exceeds 32767 characters; split Markdown content')
        c=ws.cell(row,col);c.value=value;c.data_type='s';c.font=font(role);c.alignment=Alignment(vertical='top',wrap_text=True)
        addresses = [url for _,url in url_parts(value, content.get('web_links', [])) if url]
        if len(addresses) == 1:c.hyperlink = addresses[0]
        return c
    def paragraph(ws,row,value,role='body'):
        ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=6);cell(ws,row,1,value,role)
        ws.row_dimensions[row].height=max(24, (len(str(value))//90+str(value).count('\n')+1)*20)
        return row+1
    if theme.get('cover',{}).get('enabled',True):
        ws=sheet('표지');r=paragraph(ws,1,fmt(theme.get('cover',{}).get('title','{title}')),'title')
        for value in (content['summary'],project['organization'],project['name'],version_text(content, theme),project['classification']):r=paragraph(ws,r,value)
        if content.get('example'):paragraph(ws,r,'가상 예제 데이터로 작성된 문서입니다.')
    if theme.get('copyright',{}).get('enabled',False):
        ws=sheet('저작권');paragraph(ws,1,fmt(theme['copyright'].get('title','저작권')),'heading_1');paragraph(ws,2,fmt(theme['copyright'].get('text','')))
    index=sheet(theme.get('contents',{}).get('title','목차')) if theme.get('contents',{}).get('enabled',True) else None
    if index: paragraph(index,1,content['title'],'title')
    for n,section in enumerate(content['sections'],1):
        ws=sheet(f'{n} {section["title"]}');r=paragraph(ws,1,section['title'],'heading_1')+1
        if index:
            c=cell(index,n+2,1,section['title']);c.hyperlink="#'"+ws.title.replace("'","''")+"'!A1"
        for b in section_blocks(section):
            k=b['type']
            if k=='table':
                table_start=r
                for ri,row in enumerate([b['headers'],*b['rows']]):
                    for ci,value in enumerate(row,1):
                        c=cell(ws,r,ci,value,'table_header' if ri==0 else 'table_body')
                        c.border=Border(bottom=Side(style='thin',color=theme['colors']['border'][1:]))
                        if ri==0:c.fill=PatternFill('solid',fgColor=theme['colors']['table_fill'][1:]);c.font=font('table_header',theme['colors']['table_text'][1:])
                    ws.row_dimensions[r].height=max(30,max((len(v)//18+v.count('\n')+1)*17 for v in row));r+=1
                if not ws.print_title_rows:ws.print_title_rows=f'1:{table_start}'
                r+=1
            elif k=='image':
                img=Image(b['path']);scale=min(1,900/img.width);img.width*=scale;img.height*=scale;ws.add_image(img,f'A{r}')
                r+=int(img.height/20)+2;r=paragraph(ws,r,b['alt'])
            else:
                value=(theme['notes'][b['role']]['label']+' · ' if k=='note' else '')+b['text']
                r=paragraph(ws,r,value,'heading_'+str(b['level']) if k=='heading' else 'body')
        ws.print_area=f'A1:F{max(1,r)}'
    if public_references(content):
        ws=sheet('참고 자료');r=1
        for ref in public_references(content):r=paragraph(ws,r,ref['title']+('\n'+ref['url'] if ref.get('url') else ''))
    if theme['ending'].get('enabled',True):
        ws=sheet('마지막 페이지');paragraph(ws,1,fmt(theme['ending']['title']),'title');paragraph(ws,2,fmt(theme['ending']['text']))
    wb.save(path)
