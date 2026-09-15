"""Document harness CLI. Content, contracts and visual tokens stay independent."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from tempfile import TemporaryDirectory
from uuid import uuid4
from output_adapters import style_token, public_references, render_html, render_xlsx

ROOT = Path(__file__).resolve().parent


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8-sig"))


def write_json(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def require(condition, message):
    if not condition:
        raise ValueError(message)


def text(value):
    return isinstance(value, str) and bool(value.strip())


def validate(content, contract, project, theme, release=False):
    """Validate before writing any output; never repair or invent missing content."""
    require(isinstance(content, dict), "document: object required")
    for key in ("id", "type", "title", "version", "summary"):
        require(text(content.get(key)), f"document.{key}: nonempty string required")
    require(re.fullmatch(r"[A-Za-z0-9_-]+", content["id"]), "document.id: safe ASCII identifier required")
    require(content["type"] == contract["id"], "document.type does not match template")
    for key in ("name", "organization", "contact", "classification"):
        require(text(project.get(key)), f"project.{key}: required")
    require(isinstance(content.get("sources"), list) and content["sources"], "sources required")
    source_ids = set()
    for source in content["sources"]:
        require(isinstance(source, dict), "source must be an object")
        for key in ("id", "title", "locator"):
            require(text(source.get(key)), f"source.{key}: required")
        require(source["id"] not in source_ids, "duplicate source ID")
        source_ids.add(source["id"])
    sections = content.get("sections")
    require(isinstance(sections, list) and sections, "sections required")
    ids = []
    warnings = []
    for section in sections:
        require(isinstance(section, dict), "section must be an object")
        require(text(section.get("id")) and text(section.get("title")), "section id/title required")
        ids.append(section["id"])
        require(section.get("status") in ("confirmed", "pending", "not_applicable"), "section status required")
        if section["status"] != "confirmed":
            require(text(section.get("reason")), "pending/not_applicable section needs reason")
        if section["status"] == "pending":
            require(text(section.get("owner")), "pending section needs owner")
            warnings.append(f"{section['id']}: 미확정 / {section['owner']} / {section['reason']}")
        refs = section.get("sources", [])
        require(isinstance(refs, list) and all(r in source_ids for r in refs), "unknown source reference")
        require(refs or section["status"] != "confirmed", f"{section['id']}: confirmed content needs source")
        blocks = section.get("blocks")
        require(isinstance(blocks, list), "section.blocks must be a list")
        require(blocks or section["status"] != "confirmed", "confirmed section must have content")
        for block in blocks:
            require(isinstance(block, dict), "block must be an object")
            kind = block.get("type")
            require(kind in ("paragraph", "note", "table", "code", "heading", "image"), f"unsupported block type: {kind}")
            if kind == "image":
                require(text(block.get("alt")) and Path(block.get("path", "")).is_file(), "image alt and existing file required")
                require(Path(block["path"]).suffix.lower() in (".png", ".jpg", ".jpeg"), "only PNG/JPEG images supported")
                continue
            if kind == "heading":
                require(type(block.get("level")) is int and 2 <= block["level"] <= 4, "heading level must be 2..4")
            if kind == "table":
                headers, rows = block.get("headers"), block.get("rows")
                require(isinstance(headers, list) and 1 <= len(headers) <= 6 and all(text(h) for h in headers), "table needs 1-6 named columns; split wide tables")
                require(isinstance(rows, list) and rows, "table rows required")
                require(all(isinstance(row, list) and len(row) == len(headers) and all(isinstance(v, str) for v in row) for row in rows), "table cells must be strings, matching headers")
            else:
                require(text(block.get("text")), f"{kind}.text required")
            if kind == "note":
                require(block.get("role") in ("info", "warning", "example"), "invalid semantic note role")
    require(len(ids) == len(set(ids)), "duplicate section IDs")
    missing = set(contract["required_sections"]) - set(ids)
    require(not missing, f"missing required sections: {sorted(missing)}")
    for section in sections:
        require(all(ref in ids for ref in section.get("refs", [])), "unknown section reference")
    if content.get("example", False):
        warnings.append("가상 예제 데이터입니다. 실제 제품 검증 결과가 아닙니다.")
    if release:
        require(not warnings, "release blocked: example or pending content")
    public_references(content)
    validate_theme(theme)
    return warnings


def validate_theme(t):
    for key in ("id", "font_family"):
        require(text(t.get(key)), f"theme.{key} required")
    for key in ("page_width_mm", "page_height_mm", "margin_mm", "body_pt", "title_pt", "heading_pt", "table_pt", "line_spacing"):
        require(isinstance(t.get(key), (int, float)) and 0 < t[key] < 500, f"invalid theme.{key}")
    require(t["margin_mm"] * 2 < min(t["page_width_mm"], t["page_height_mm"]) - 40, "theme margins leave no usable space")
    for key in ("text", "muted", "accent", "table_fill", "table_text", "border", "background"):
        require(re.fullmatch(r"#[0-9A-Fa-f]{6}", t.get("colors", {}).get(key, "")), f"invalid color: {key}")
    for role in ("info", "warning", "example"):
        config = t.get("notes", {}).get(role, {})
        require(text(config.get("label")), f"note label required: {role}")
        for field in ("fill", "marker"):
            require(re.fullmatch(r"#[0-9A-Fa-f]{6}", config.get(field, "")), "invalid note color")
    for group, fields in {"header": ("left", "right"), "footer": ("left",), "ending": ("title", "text")}.items():
        require(isinstance(t.get(group), dict), f"theme.{group} required")
        for field in fields:
            require(text(t[group].get(field)), f"theme.{group}.{field} required")
    allowed = {"title", "version", "id", "name", "organization", "contact", "classification"}
    for group in ("header", "footer", "ending"):
        for value in t[group].values():
            if isinstance(value, str):
                require(set(re.findall(r"{([^}]+)}", value)) <= allowed, "unknown theme placeholder")
    require(not re.search(r'[<>;{}\r\n]', t["font_family"]), "invalid font family")
    for group in ("cover", "copyright", "contents"):
        cfg = t.get(group, {})
        require(isinstance(cfg, dict), f"theme.{group}: object required")
        require(type(cfg.get("enabled", True)) is bool, f"theme.{group}.enabled must be boolean")
        for key in ("title", "text"):
            if key in cfg:
                require(isinstance(cfg[key], str), f"theme.{group}.{key}: string required")
                require(set(re.findall(r"{([^}]+)}", cfg[key])) <= allowed, "unknown theme placeholder")
    require(isinstance(t.get("styles", {}), dict), "theme.styles must be object")
    for role, override in t.get("styles", {}).items():
        require(role in ("title", "heading_1", "heading_2", "heading_3", "heading_4", "body", "table_header", "table_body", "header", "footer"), "unknown style role")
        require(isinstance(override, dict) and set(override) <= {"font_family", "size_pt", "bold"}, "invalid style override")
        token = style_token(t, role)
        require(text(token["font_family"]) and not re.search(r'[<>;{}\r\n]', token["font_family"]), "invalid style font family")
        require(type(token["size_pt"]) in (int, float) and 0 < token["size_pt"] < 100, "invalid style size")
        require(type(token["bold"]) is bool, "style bold must be boolean")
    ts = t.get("table", {})
    for key, default in (("padding_x_pt", 7), ("padding_y_pt", 7), ("border_pt", .35)):
        value = ts.get(key, default)
        require(type(value) in (int, float) and 0 <= value <= 24, f"invalid table.{key}")


def interpolate(value, content, project):
    # Legacy company templates may contain internal identifiers; never publish those lines.
    value = "\n".join(line for line in value.split("\n") if "{id}" not in line)
    return value.format_map({**project, **content})


def section_blocks(section):
    if section["status"] != "confirmed":
        label = "미확정" if section["status"] == "pending" else "해당 없음"
        yield {"type": "note", "role": "warning", "text": f"{label}: {section['reason']}"}
    yield from section["blocks"]



def render_docx(path, content, project, theme):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    d = Document()
    table_tokens = theme.get("table", {})
    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(theme["page_width_mm"]), Mm(theme["page_height_mm"])
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Mm(theme["margin_mm"])
    sec.header_distance = sec.footer_distance = Mm(10)
    sec.different_first_page_header_footer = True
    for name, size in (("Normal", theme["body_pt"]), ("Title", theme["title_pt"]), ("Heading 1", theme["heading_pt"]), ("Heading 2", theme["body_pt"] + 2), ("Heading 3", theme["body_pt"] + 1), ("Heading 4", theme["body_pt"]), ("Header", 9), ("Footer", 9)):
        style = d.styles[name]
        key = {"Normal": "body", "Title": "title", "Header": "header", "Footer": "footer"}.get(name, name.lower().replace(" ", "_"))
        token = style_token(theme, key)
        style.font.name = token["font_family"]
        style.font.bold = token["bold"]
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), token["font_family"])
        style.font.size = Pt(token["size_pt"])
        style.font.color.rgb = RGBColor.from_string("000000" if name != "Normal" else theme["colors"]["text"][1:])
        style.paragraph_format.line_spacing = theme["line_spacing"]
        style.paragraph_format.space_after = Pt(8)
        if name.startswith("Heading"):
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_before = Pt(15)
    hp = sec.header.paragraphs[0]
    hp.text = interpolate(theme["header"]["left"], content, project) + "  |  " + interpolate(theme["header"]["right"], content, project)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.add_run(interpolate(theme["footer"]["left"], content, project) + "  ·  ")
    for instr in ("PAGE", "NUMPAGES"):
        if instr == "NUMPAGES":
            fp.add_run(" / ")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), instr)
        fp._p.append(field)
    d.core_properties.title = content["title"]
    d.core_properties.author = project["organization"]
    d.core_properties.comments = ""
    if theme.get("cover", {}).get("enabled", True):
        d.add_paragraph(project["organization"])
        d.add_paragraph(interpolate(theme.get("cover", {}).get("title", "{title}"), content, project), "Title")
        d.add_paragraph(content["summary"])
        d.add_paragraph(f"{content['version']}\n{project['name']}\n{project['classification']}")
        if content.get("example"):
            d.add_paragraph("가상 예제 데이터로 작성된 문서입니다.")
        d.add_page_break()
    if theme.get("copyright", {}).get("enabled", False):
        d.add_heading(interpolate(theme["copyright"].get("title", "저작권"), content, project), 1)
        d.add_paragraph(interpolate(theme["copyright"].get("text", ""), content, project))
        d.add_page_break()
    if theme.get("contents", {}).get("enabled", True):
        d.add_heading(theme.get("contents", {}).get("title", "문서 구성"), 1)
        for n, section in enumerate(content["sections"], 1):
            d.add_paragraph(f"{n}. {section['title']}")
        d.add_page_break()
    for n, section in enumerate(content["sections"], 1):
        d.add_heading(f"{n} {section['title']}", 1)
        for block in section_blocks(section):
            kind = block["type"]
            if kind == "heading":
                d.add_heading(block["text"], block["level"])
            elif kind == "image":
                from PIL import Image as PILImage
                with PILImage.open(block["path"]) as img:
                    w, h = img.size
                max_w = sec.page_width - sec.left_margin - sec.right_margin
                max_h = sec.page_height - sec.top_margin - sec.bottom_margin - Mm(20)
                scale = min(max_w / w, max_h / h)
                picture = d.add_paragraph().add_run().add_picture(block["path"], width=int(w * scale), height=int(h * scale))
                picture._inline.docPr.set("descr", block["alt"])
                d.add_paragraph(block["alt"])
            elif kind == "table":
                table = d.add_table(rows=1, cols=len(block["headers"]))
                table.autofit = False
                width = (sec.page_width - sec.left_margin - sec.right_margin) / len(block["headers"])
                for col in table.columns:
                    col.width = int(width)
                for idx, value in enumerate(block["headers"]):
                    table.rows[0].cells[idx].text = value
                repeat = OxmlElement("w:tblHeader")
                table.rows[0]._tr.get_or_add_trPr().append(repeat)
                for row in block["rows"]:
                    cells = table.add_row().cells
                    for idx, value in enumerate(row):
                        cells[idx].text = value
                for ri, row in enumerate(table.rows):
                    # Allow long rows to continue to the next page; repeat column labels.
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            p.paragraph_format.space_after = Pt(5)
                            for run in p.runs:
                                token = style_token(theme, "table_header" if ri == 0 else "table_body")
                                run.font.size = Pt(token["size_pt"])
                                run.font.name = token["font_family"]
                                run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), token["font_family"])
                                run.bold = token["bold"]
                                if ri == 0:
                                    run.font.color.rgb = RGBColor.from_string(theme["colors"]["table_text"][1:])
                        tc_props = cell._tc.get_or_add_tcPr()
                        margins = OxmlElement("w:tcMar")
                        for side, points in {"left": table_tokens.get("padding_x_pt", 7), "right": table_tokens.get("padding_x_pt", 7), "top": table_tokens.get("padding_y_pt", 7), "bottom": table_tokens.get("padding_y_pt", 7)}.items():
                            item = OxmlElement("w:" + side)
                            item.set(qn("w:w"), str(int(points * 20)))
                            item.set(qn("w:type"), "dxa")
                            margins.append(item)
                        tc_props.append(margins)
                        borders = OxmlElement("w:tcBorders")
                        edge = OxmlElement("w:bottom")
                        for key, val in {"val": "single", "sz": str(int(table_tokens.get("border_pt", .35) * 8)), "color": theme["colors"]["border"][1:]}.items():
                            edge.set(qn("w:" + key), val)
                        borders.append(edge)
                        tc_props.append(borders)
                        if ri == 0:
                            shading = OxmlElement("w:shd")
                            shading.set(qn("w:fill"), theme["colors"]["table_fill"][1:])
                            cell._tc.get_or_add_tcPr().append(shading)
                d.add_paragraph()
            else:
                value = block["text"]
                if kind == "note":
                    note = theme["notes"][block["role"]]
                    p = d.add_paragraph()
                    p.add_run(note["label"] + "  ").bold = True
                    p.add_run(value)
                    props = p._p.get_or_add_pPr()
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), note["fill"][1:])
                    props.append(shade)
                    border = OxmlElement("w:pBdr")
                    left = OxmlElement("w:left")
                    for key, val in {"val": "single", "sz": "18", "space": "8", "color": note["marker"][1:]}.items():
                        left.set(qn("w:" + key), val)
                    border.append(left)
                    props.append(border)
                else:
                    d.add_paragraph(value)
    if public_references(content):
        d.add_heading("참고 자료", 1)
        for source in public_references(content):
            d.add_paragraph(source["title"] + ("\n" + source["url"] if source.get("url") else ""))
    if theme["ending"].get("enabled", True):
        d.add_page_break()
        d.add_heading(interpolate(theme["ending"]["title"], content, project), 1)
        d.add_paragraph(interpolate(theme["ending"]["text"], content, project))
    d.save(path)


def render_pdf(path, content, project, theme, runtime):
    from html import escape
    from reportlab.lib.colors import HexColor
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle, Image

    require(Path(runtime.get("font_regular", "")).is_file(), "configure runtime font_regular path")
    require(Path(runtime.get("font_bold", "")).is_file(), "configure runtime font_bold path")
    pdfmetrics.registerFont(TTFont("BodyFont", runtime["font_regular"]))
    pdfmetrics.registerFont(TTFont("BoldFont", runtime["font_bold"]))
    pdfmetrics.registerFontFamily("BodyFont", normal="BodyFont", bold="BoldFont")
    body = ParagraphStyle("body", fontName="BodyFont", fontSize=theme["body_pt"], leading=theme["body_pt"] * theme["line_spacing"], spaceAfter=9, wordWrap="CJK", textColor=HexColor(theme["colors"]["text"]))
    heading = ParagraphStyle("heading", parent=body, fontName="BoldFont", fontSize=theme["heading_pt"], leading=theme["heading_pt"] * 1.4, spaceBefore=16, spaceAfter=10, keepWithNext=True)
    title = ParagraphStyle("title", parent=heading, fontSize=theme["title_pt"], leading=theme["title_pt"] * 1.4, textColor=HexColor("#000000"))
    cell_style = ParagraphStyle("cell", parent=body, fontSize=theme["table_pt"], leading=theme["table_pt"] * 1.45, spaceAfter=0)
    head_cell = ParagraphStyle("headcell", parent=cell_style, fontName="BoldFont", textColor=HexColor(theme["colors"]["table_text"]))
    def apply_token(style, role):
        token = style_token(theme, role)
        style.fontName = "BoldFont" if token["bold"] else "BodyFont"
        style.fontSize = token["size_pt"]
        style.leading = token["size_pt"] * theme["line_spacing"]
        return style
    for style, role in ((body, "body"), (heading, "heading_1"), (title, "title"), (cell_style, "table_body"), (head_cell, "table_header")):
        apply_token(style, role)
    width, height = theme["page_width_mm"] * mm, theme["page_height_mm"] * mm
    margin = theme["margin_mm"] * mm
    usable = width - 2 * margin
    table_tokens = theme.get("table", {})

    def para(s, style=body):
        return Paragraph(escape(s).replace("\n", "<br/>"), style)

    def decorate(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(HexColor(theme["colors"]["background"]))
        canvas.rect(0, 0, width, height, fill=1, stroke=0)
        if doc.page > 1:
            token = style_token(theme, "header")
            canvas.setFont("BoldFont" if token["bold"] else "BodyFont", token["size_pt"])
            canvas.setFillColor(HexColor(theme["colors"]["muted"]))
            canvas.drawString(margin, height - 12 * mm, interpolate(theme["header"]["left"], content, project))
            canvas.drawRightString(width - margin, height - 12 * mm, interpolate(theme["header"]["right"], content, project))
            token = style_token(theme, "footer")
            canvas.setFont("BoldFont" if token["bold"] else "BodyFont", token["size_pt"])
            canvas.drawString(margin, 12 * mm, interpolate(theme["footer"]["left"], content, project))
            canvas.drawRightString(width - margin, 12 * mm, str(doc.page))
        canvas.restoreState()

    story = [Spacer(1, 25 * mm), para(project["organization"]), para(interpolate(theme.get("cover", {}).get("title", "{title}"), content, project), title), para(content["summary"]), Spacer(1, 10 * mm), para(content["version"]), para(project["name"]), para(project["classification"])]
    if content.get("example"):
        story.append(para("가상 예제 데이터로 작성된 문서입니다."))
    if not theme.get("cover", {}).get("enabled", True):
        story = []
    if theme.get("copyright", {}).get("enabled", False):
        if story:
            story.append(PageBreak())
        story.extend([para(interpolate(theme["copyright"].get("title", "저작권"), content, project), heading), para(interpolate(theme["copyright"].get("text", ""), content, project))])
    if theme.get("contents", {}).get("enabled", True):
        if story:
            story.append(PageBreak())
        story.append(para(theme.get("contents", {}).get("title", "문서 구성"), heading))
        for n, section in enumerate(content["sections"], 1):
            story.append(para(f"{n}. {section['title']}"))
    if story:
        story.append(PageBreak())
    for n, section in enumerate(content["sections"], 1):
        story.append(para(f"{n} {section['title']}", heading))
        for block in section_blocks(section):
            if block["type"] == "heading":
                size = theme["body_pt"] + (4 - block["level"])
                subheading = ParagraphStyle("h" + str(block["level"]), parent=heading, fontSize=size, leading=size*1.5)
                apply_token(subheading, "heading_" + str(block["level"]))
                story.append(para(block["text"], subheading))
            elif block["type"] == "image":
                img = Image(block["path"], hAlign="LEFT")
                scale = min(usable / img.imageWidth, (height-2*margin-40)/img.imageHeight)
                img.drawWidth = img.imageWidth * scale
                img.drawHeight = img.imageHeight * scale
                story.extend([img, para(block["alt"])])
            elif block["type"] == "table":
                rows = [[para(h, head_cell) for h in block["headers"]]]
                rows.extend([[para(v, cell_style) for v in row] for row in block["rows"]])
                tab = Table(rows, colWidths=[usable / len(block["headers"])] * len(block["headers"]), repeatRows=1, hAlign="LEFT", splitByRow=1, splitInRow=1)
                tab.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), HexColor(theme["colors"]["table_fill"])), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEBELOW", (0, 0), (-1, -1), table_tokens.get("border_pt", .35), HexColor(theme["colors"]["border"])), ("LEFTPADDING", (0, 0), (-1, -1), table_tokens.get("padding_x_pt", 7)), ("RIGHTPADDING", (0, 0), (-1, -1), table_tokens.get("padding_x_pt", 7)), ("TOPPADDING", (0, 0), (-1, -1), table_tokens.get("padding_y_pt", 7)), ("BOTTOMPADDING", (0, 0), (-1, -1), table_tokens.get("padding_y_pt", 7))]))
                story.extend([tab, Spacer(1, 8)])
            elif block["type"] == "note":
                note = theme["notes"][block["role"]]
                tab = Table([[para(note["label"] + "  " + block["text"])]], colWidths=[usable], splitInRow=1)
                tab.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), HexColor(note["fill"])), ("LINEBEFORE", (0, 0), (0, -1), 3, HexColor(note["marker"])), ("LEFTPADDING", (0, 0), (-1, -1), 10), ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
                story.extend([tab, Spacer(1, 9)])
            else:
                story.append(para(block["text"]))
    if public_references(content):
        story.append(para("참고 자료", heading))
        for source in public_references(content):
            story.append(para(source["title"] + ("\n" + source["url"] if source.get("url") else "")))
    if theme["ending"].get("enabled", True):
        story.extend([PageBreak(), Spacer(1, 20 * mm), para(interpolate(theme["ending"]["title"], content, project), title), para(interpolate(theme["ending"]["text"], content, project))])
    doc = SimpleDocTemplate(str(path), pagesize=(width, height), leftMargin=margin, rightMargin=margin, topMargin=margin, bottomMargin=margin, title=content["title"], author=project["organization"])
    doc.build(story, onFirstPage=decorate, onLaterPages=decorate)


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def build(content_path, template_path, project_path, theme_path, runtime_path, output, formats, release=False):
    paths = [Path(p).resolve() for p in (content_path, template_path, project_path, theme_path)]
    content, contract, project, theme = map(read_json, paths)
    runtime = read_json(runtime_path) if runtime_path else {}
    warnings = validate(content, contract, project, theme, release)
    require(formats and set(formats) <= {"docx", "pdf", "html", "xlsx"}, "formats must be docx/pdf/html/xlsx")
    if "docx" in formats:
        require(theme["colors"]["background"] == "#FFFFFF", "DOCX page background currently supports white only; use PDF for colored background")
    output = Path(output).resolve()
    output.mkdir(parents=True, exist_ok=True)
    run = content["id"] + "-" + datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ") + "-" + uuid4().hex[:8]
    dest = output / run
    # Failures never publish partially built runs or overwrite a previous output.
    with TemporaryDirectory(prefix=".build-", dir=output) as tmp:
        staging = Path(tmp)
        for fmt in formats:
            path = staging / ("document." + fmt)
            if fmt == "docx":
                render_docx(path, content, project, theme)
            elif fmt == "pdf":
                render_pdf(path, content, project, theme, runtime)
            elif fmt == "html":
                render_html(path, content, project, theme)
            else:
                render_xlsx(path, content, project, theme)
        manifest = {"run_id": run, "created_at": datetime.now(timezone.utc).isoformat(), "template": contract["id"], "inputs": [{"path": str(p), "sha256": digest(p)} for p in paths], "outputs": {p.name: digest(p) for p in staging.iterdir()}, "validation": "passed", "warnings": warnings, "visual_review": {fmt: "not_reviewed" for fmt in formats}, "pdf_origin": "independent ReportLab render; not a DOCX conversion" if "pdf" in formats else None}
        write_json(staging / "manifest.json", manifest)
        write_json(staging / "resolved-input.json", {"document": content, "template": contract, "project": project, "theme": theme})
        staging.rename(dest)
    return dest


def main():
    parser = argparse.ArgumentParser(description="Document Creator: validate / build")
    parser.add_argument("command", choices=["validate", "build"])
    parser.add_argument("--input", required=True)
    parser.add_argument("--template", required=True)
    parser.add_argument("--project", required=True)
    parser.add_argument("--theme", default=str(ROOT / "themes/blue-office.json"))
    parser.add_argument("--runtime", default=str(ROOT / ".local/runtime.json"))
    parser.add_argument("--output", default=str(ROOT / "output"))
    parser.add_argument("--formats", nargs="+", default=["pdf"], choices=["pdf", "docx", "html", "xlsx"])
    parser.add_argument("--release", action="store_true", help="Reject example/pending data; does not imply visual QA approval")
    args = parser.parse_args()
    try:
        if args.command == "validate":
            warnings = validate(read_json(args.input), read_json(args.template), read_json(args.project), read_json(args.theme), args.release)
            print(json.dumps({"validation": "passed", "warnings": warnings}, ensure_ascii=False, indent=2))
        else:
            print(build(args.input, args.template, args.project, args.theme, args.runtime, args.output, args.formats, args.release))
        return 0
    except (ValueError, KeyError, TypeError, OSError, ImportError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
